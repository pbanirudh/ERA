import pandas as pd
import numpy as np
import datetime
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def generate_retail_data():
    print("Generating mock e-commerce transactions...")
    np.random.seed(42)
    
    # Setup configuration
    n_transactions = 1000
    n_customers = 80
    
    # Lists for random draws
    countries = ['United Kingdom'] * 800 + ['Germany'] * 50 + ['France'] * 50 + ['Spain'] * 40 + ['EIRE'] * 30 + ['Netherlands'] * 30
    products = [
        ('85123A', 'WHITE HANGING HEART T-LIGHT HOLDER', 2.55),
        ('22423', 'REGENCY CAKESTAND 3 TIER', 12.75),
        ('84879', 'ASSORTED COLOUR BIRD ORNAMENT', 1.69),
        ('47566', 'PARTY BUNTING', 4.95),
        ('22720', 'SET OF 3 CAKE TINS PANTRY DESIGN', 4.95),
        ('20725', 'LUNCH BAG RED RETROSPOT', 1.65),
        ('23084', 'RABBIT NIGHT LIGHT', 2.08),
        ('22197', 'POPCORN HOLDER', 0.85),
        ('22383', 'LUNCH BAG SUKI DESIGN', 1.65),
        ('85099B', 'JUMBO BAG RED RETROSPOT', 2.08),
    ]
    
    customer_ids = [str(12000 + i) for i in range(1, n_customers + 1)]
    
    # Generate transactional records
    records = []
    start_date = datetime.datetime(2010, 12, 1, 8, 0)
    
    # Store customer first purchase to calculate realistic cohorts later
    customer_acquisition_dates = {}
    
    for i in range(n_transactions):
        invoice_no = str(536365 + np.random.randint(0, 1000))
        cust_id = np.random.choice(customer_ids)
        
        # Decide invoice date. Customers make purchases over time
        # Add random number of days and minutes to the start_date
        days_offset = np.random.randint(0, 373) # 1 year
        minutes_offset = np.random.randint(0, 1440)
        invoice_date = start_date + datetime.timedelta(days=days_offset, minutes=minutes_offset)
        
        # Track acquisition date
        if cust_id not in customer_acquisition_dates or invoice_date < customer_acquisition_dates[cust_id]:
            customer_acquisition_dates[cust_id] = invoice_date
            
        prod = products[np.random.randint(0, len(products))]
        quantity = np.random.randint(1, 48)
        unit_price = prod[2]
        total_amount = round(quantity * unit_price, 2)
        country = np.random.choice(countries)
        
        records.append({
            'InvoiceNo': invoice_no,
            'StockCode': prod[0],
            'Description': prod[1],
            'Quantity': quantity,
            'InvoiceDate': invoice_date,
            'UnitPrice': unit_price,
            'CustomerID': cust_id,
            'Country': country,
            'TotalAmount': total_amount
        })
        
    df_tx = pd.DataFrame(records).sort_values(by='InvoiceDate').reset_index(drop=True)
    
    # --- Generate RFM Analysis ---
    print("Calculating RFM metrics...")
    max_date = df_tx['InvoiceDate'].max()
    
    rfm_records = []
    for cust_id in customer_ids:
        cust_tx = df_tx[df_tx['CustomerID'] == cust_id]
        if cust_tx.empty:
            continue
            
        last_date = cust_tx['InvoiceDate'].max()
        recency = (max_date - last_date).days
        frequency = cust_tx['InvoiceNo'].nunique()
        monetary = cust_tx['TotalAmount'].sum()
        
        rfm_records.append({
            'CustomerID': cust_id,
            'Recency': recency,
            'Frequency': frequency,
            'MonetaryValue': round(monetary, 2)
        })
        
    df_rfm = pd.DataFrame(rfm_records)
    
    # Calculate segments
    spend_ranks = df_rfm['MonetaryValue'].rank(pct=True, ascending=False)
    df_rfm['CustomerSegment'] = np.where(
        spend_ranks <= 0.10, 'Top 10% (Platinum Tier)',
        np.where(spend_ranks <= 0.25, 'Top 25% (Gold Tier)', 'General Tier')
    )
    df_rfm = df_rfm.sort_values(by='MonetaryValue', ascending=False).reset_index(drop=True)
    
    # --- Generate Cohort Data ---
    print("Simulating cohort retention matrix...")
    df_tx['AcquisitionMonth'] = df_tx['CustomerID'].map(lambda cid: customer_acquisition_dates[cid].replace(day=1, hour=0, minute=0, second=0))
    df_tx['ActivityMonth'] = df_tx['InvoiceDate'].map(lambda d: d.replace(day=1, hour=0, minute=0, second=0))
    
    # Calculate months elapsed
    df_tx['CohortIndex'] = ((df_tx['ActivityMonth'].dt.year - df_tx['AcquisitionMonth'].dt.year) * 12 + 
                            (df_tx['ActivityMonth'].dt.month - df_tx['AcquisitionMonth'].dt.month))
    
    # Calculate cohort size
    cohort_sizes = df_tx.groupby('AcquisitionMonth')['CustomerID'].nunique().to_dict()
    
    cohort_groups = df_tx.groupby(['AcquisitionMonth', 'CohortIndex'])
    cohort_data = []
    
    for (acq_month, index), group in cohort_groups:
        retained = group['CustomerID'].nunique()
        size = cohort_sizes[acq_month]
        ret_rate = round((retained / size), 4)
        revenue = group['TotalAmount'].sum()
        
        cohort_data.append({
            'AcquisitionMonth': acq_month.strftime('%Y-%m'),
            'CohortSize': size,
            'CohortIndex': int(index),
            'RetainedCustomers': retained,
            'RetentionRatePercent': ret_rate,
            'CohortRevenue': round(revenue, 2)
        })
        
    df_cohort = pd.DataFrame(cohort_data).sort_values(by=['AcquisitionMonth', 'CohortIndex']).reset_index(drop=True)
    
    # Ensure export path
    os.makedirs('excel', exist_ok=True)
    xlsx_path = os.path.join('excel', 'Online_Retail_Analytical_Model.xlsx')
    
    # Save sheets
    print("Writing to Excel workbook and applying premium styles...")
    with pd.ExcelWriter(xlsx_path, engine='openpyxl') as writer:
        df_tx.to_excel(writer, sheet_name='Transactions', index=False)
        df_rfm.to_excel(writer, sheet_name='RFM_Analysis', index=False)
        df_cohort.to_excel(writer, sheet_name='Cohort_Data', index=False)
        
    # Openpyxl styling
    wb = openpyxl.load_workbook(xlsx_path)
    
    # Styles
    font_family = "Segoe UI"
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid") # Deep navy
    header_font = Font(name=font_family, size=11, bold=True, color="FFFFFF")
    data_font = Font(name=font_family, size=10)
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )
    
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        ws.views.sheetView[0].showGridLines = True
        
        # Header formatting
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=False)
            ws.row_dimensions[1].height = 28
            
        # Data formatting
        max_row = ws.max_row
        max_col = ws.max_column
        
        for row in range(2, max_row + 1):
            ws.row_dimensions[row].height = 20
            for col in range(1, max_col + 1):
                cell = ws.cell(row=row, column=col)
                cell.font = data_font
                cell.border = thin_border
                
        # Column specific formats
        if sheet_name == 'Transactions':
            # Quantity (Col D): Int
            # InvoiceDate (Col E): Date
            # UnitPrice (Col F): Currency
            # TotalAmount (Col I): Currency
            for row in range(2, max_row + 1):
                ws.cell(row=row, column=4).number_format = '#,##0'
                ws.cell(row=row, column=5).number_format = 'yyyy-mm-dd hh:mm'
                ws.cell(row=row, column=6).number_format = '$#,##0.00'
                ws.cell(row=row, column=9).number_format = '$#,##0.00'
                
        elif sheet_name == 'RFM_Analysis':
            # Recency (Col B): Int
            # Frequency (Col C): Int
            # MonetaryValue (Col D): Currency
            for row in range(2, max_row + 1):
                ws.cell(row=row, column=2).number_format = '#,##0'
                ws.cell(row=row, column=3).number_format = '#,##0'
                ws.cell(row=row, column=4).number_format = '$#,##0.00'
                
        elif sheet_name == 'Cohort_Data':
            # CohortSize (Col B): Int
            # CohortIndex (Col C): Int
            # RetainedCustomers (Col D): Int
            # RetentionRatePercent (Col E): Percent
            # CohortRevenue (Col F): Currency
            for row in range(2, max_row + 1):
                ws.cell(row=row, column=2).number_format = '#,##0'
                ws.cell(row=row, column=3).number_format = '0'
                ws.cell(row=row, column=4).number_format = '#,##0'
                ws.cell(row=row, column=5).number_format = '0.00%'
                ws.cell(row=row, column=6).number_format = '$#,##0.00'

        # Auto fit column widths
        for col in range(1, max_col + 1):
            col_letter = get_column_letter(col)
            max_len = 0
            for row in range(1, max_row + 1):
                val = str(ws.cell(row=row, column=col).value or '')
                if len(val) > max_len:
                    max_len = len(val)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)
            
    wb.save(xlsx_path)
    print(f"Workbook successfully saved to {xlsx_path}!")

if __name__ == '__main__':
    generate_retail_data()
